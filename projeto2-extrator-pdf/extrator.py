"""
Extrator de dados de PDFs para Excel.

O que faz:
    1. Abre todos os PDFs de uma pasta.
    2. Le o texto de cada um.
    3. Procura os campos: numero da nota, data, cliente, CNPJ e valor total.
    4. Junta tudo numa planilha Excel (uma linha por PDF).
    5. Se algum campo nao for encontrado, marca o PDF como "CONFERIR".

Como rodar:
    python extrator.py pdfs_exemplo
Resultado: saida/notas_extraidas.xlsx

Tudo roda no seu computador: nenhum arquivo e enviado para a internet.
"""
import argparse
import re
import sys
from pathlib import Path

import pandas as pd
import pdfplumber
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

# Cada campo tem um "padrao de busca" (expressao regular).
# O que esta entre parenteses ( ) e o pedaco que queremos guardar.
CAMPOS = {
    "Numero": r"Nota N[oº°]?:?\s*(\d+)",
    "Data": r"Data de emiss[aã]o:\s*(\d{2}/\d{2}/\d{4})",
    "Cliente": r"Cliente:\s*(.+)",
    "CNPJ": r"CNPJ do cliente:\s*(\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2})",
    "Valor Total": r"VALOR TOTAL:\s*R\$\s*([\d.,]+)",
}


def ler_texto_pdf(arquivo: Path) -> str:
    """Junta o texto de todas as paginas do PDF."""
    with pdfplumber.open(arquivo) as pdf:
        return "\n".join(pagina.extract_text() or "" for pagina in pdf.pages)


def extrair_campos(texto: str) -> dict:
    """Procura cada campo no texto. Se nao achar, deixa vazio."""
    dados = {}
    for campo, padrao in CAMPOS.items():
        achado = re.search(padrao, texto, flags=re.IGNORECASE)
        dados[campo] = achado.group(1).strip() if achado else None
    # "1.234,50" (formato brasileiro) -> 1234.50 (numero que o Excel entende)
    if dados["Valor Total"]:
        dados["Valor Total"] = float(dados["Valor Total"].replace(".", "").replace(",", "."))
    return dados


def processar_pasta(pasta: Path) -> pd.DataFrame:
    pdfs = sorted(pasta.glob("*.pdf"))
    if not pdfs:
        sys.exit(f"Erro: nenhum PDF encontrado em {pasta}")

    linhas = []
    for arquivo in pdfs:
        try:
            dados = extrair_campos(ler_texto_pdf(arquivo))
            faltando = [c for c, v in dados.items() if v is None]
            dados["Status"] = "OK" if not faltando else "CONFERIR: falta " + ", ".join(faltando)
        except Exception as erro:  # PDF corrompido, protegido por senha etc.
            dados = {c: None for c in CAMPOS}
            dados["Status"] = f"ERRO ao abrir: {erro.__class__.__name__}"
        dados["Arquivo"] = arquivo.name
        linhas.append(dados)
        print(f"{arquivo.name}: {dados['Status']}")

    return pd.DataFrame(linhas, columns=["Arquivo", *CAMPOS, "Status"])


def salvar_excel(tabela: pd.DataFrame, arquivo_xlsx: Path) -> None:
    tabela.to_excel(arquivo_xlsx, index=False, sheet_name="Notas")
    planilha = load_workbook(arquivo_xlsx)
    aba = planilha["Notas"]
    for celula in aba[1]:  # cabecalho
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="1F4E78")
    amarelo = PatternFill("solid", fgColor="FFF2CC")
    for linha in aba.iter_rows(min_row=2):
        linha[5].number_format = '"R$" #,##0.00'  # coluna Valor Total
        if linha[6].value != "OK":  # destaca o que precisa de conferencia
            for celula in linha:
                celula.fill = amarelo
    for coluna in aba.columns:
        maior = max(len(str(c.value)) for c in coluna if c.value is not None)
        aba.column_dimensions[coluna[0].column_letter].width = maior + 3
    planilha.save(arquivo_xlsx)


def main():
    parser = argparse.ArgumentParser(description="Extrai dados de PDFs de notas para uma planilha Excel.")
    parser.add_argument("pasta", type=Path, help="pasta com os PDFs")
    parser.add_argument("--saida", type=Path, default=Path("saida"), help="pasta do resultado (padrao: saida)")
    args = parser.parse_args()

    args.saida.mkdir(exist_ok=True)
    arquivo_xlsx = args.saida / "notas_extraidas.xlsx"
    tabela = processar_pasta(args.pasta)
    salvar_excel(tabela, arquivo_xlsx)

    ok = (tabela["Status"] == "OK").sum()
    print(f"\nPronto! {ok} de {len(tabela)} PDFs lidos sem problema.")
    print(f"Planilha: {arquivo_xlsx}")


if __name__ == "__main__":
    main()
