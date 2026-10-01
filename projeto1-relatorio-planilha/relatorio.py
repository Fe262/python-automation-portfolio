"""
Relatorio automatico de vendas.

O que faz:
    1. Le uma planilha de vendas (Excel).
    2. Calcula totais: por mes, por produto e por vendedor.
    3. Gera um grafico do faturamento por mes.
    4. Salva tudo num Excel novo, formatado e pronto para enviar ao chefe/cliente.

Como rodar:
    python relatorio.py vendas_exemplo.xlsx
Resultado: saida/relatorio_vendas.xlsx e saida/grafico_mensal.png

A planilha de entrada precisa ter estas colunas:
    Data, Vendedor, Produto, Quantidade, Preco Unitario
"""
import argparse
import sys
from pathlib import Path

import matplotlib
matplotlib.use("Agg")  # gera o grafico direto em arquivo, sem abrir janela
import matplotlib.pyplot as plt
import pandas as pd
from openpyxl import load_workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font, PatternFill

COLUNAS_OBRIGATORIAS = ["Data", "Vendedor", "Produto", "Quantidade", "Preco Unitario"]
FORMATO_REAIS = '"R$" #,##0.00'
COR_CABECALHO = "1F4E78"


# ---------- 1. Ler e conferir a planilha ----------

def ler_vendas(arquivo: Path) -> pd.DataFrame:
    """Le a planilha e confere se tem as colunas necessarias."""
    if not arquivo.exists():
        sys.exit(f"Erro: arquivo nao encontrado: {arquivo}")

    vendas = pd.read_excel(arquivo)

    faltando = [c for c in COLUNAS_OBRIGATORIAS if c not in vendas.columns]
    if faltando:
        sys.exit(f"Erro: a planilha nao tem estas colunas: {', '.join(faltando)}")

    # Garante que cada coluna tem o tipo certo (data e numeros)
    vendas["Data"] = pd.to_datetime(vendas["Data"], errors="coerce")
    vendas["Quantidade"] = pd.to_numeric(vendas["Quantidade"], errors="coerce")
    vendas["Preco Unitario"] = pd.to_numeric(vendas["Preco Unitario"], errors="coerce")

    # Linhas com data ou numero invalido sao ignoradas (e avisamos quantas foram)
    invalidas = vendas[COLUNAS_OBRIGATORIAS].isna().any(axis=1)
    if invalidas.any():
        print(f"Aviso: {invalidas.sum()} linha(s) com dados invalidos foram ignoradas.")
    vendas = vendas[~invalidas].copy()
    if vendas.empty:
        sys.exit("Erro: nenhuma venda valida encontrada na planilha.")

    vendas["Total"] = vendas["Quantidade"] * vendas["Preco Unitario"]
    vendas["Mes"] = vendas["Data"].dt.to_period("M").astype(str)  # ex: "2026-03"
    return vendas


# ---------- 2. Calcular os resumos ----------

def calcular_resumos(vendas: pd.DataFrame) -> dict:
    """Agrupa as vendas e devolve as tabelas do relatorio."""
    por_mes = vendas.groupby("Mes", as_index=False).agg(
        Faturamento=("Total", "sum"), Vendas=("Total", "count"))

    por_produto = (vendas.groupby("Produto", as_index=False)
                   .agg(Faturamento=("Total", "sum"), Unidades=("Quantidade", "sum"))
                   .sort_values("Faturamento", ascending=False))

    por_vendedor = (vendas.groupby("Vendedor", as_index=False)
                    .agg(Faturamento=("Total", "sum"), Vendas=("Total", "count"))
                    .sort_values("Faturamento", ascending=False))

    total = vendas["Total"].sum()
    indicadores = pd.DataFrame([
        ("Faturamento total", total),
        ("Numero de vendas", len(vendas)),
        ("Ticket medio (valor medio por venda)", total / len(vendas)),
        ("Produto que mais faturou", por_produto.iloc[0]["Produto"]),
        ("Melhor vendedor", por_vendedor.iloc[0]["Vendedor"]),
        ("Periodo", f"{vendas['Data'].min():%d/%m/%Y} a {vendas['Data'].max():%d/%m/%Y}"),
    ], columns=["Indicador", "Valor"])

    return {"Resumo": indicadores, "Por mes": por_mes,
            "Por produto": por_produto, "Por vendedor": por_vendedor}


# ---------- 3. Grafico ----------

def gerar_grafico(por_mes: pd.DataFrame, arquivo_png: Path) -> None:
    """Grafico de barras do faturamento por mes."""
    fig, ax = plt.subplots(figsize=(9, 4.5))
    barras = ax.bar(por_mes["Mes"], por_mes["Faturamento"], color="#2E75B6")
    ax.set_title("Faturamento por mes", fontsize=14, fontweight="bold")
    ax.set_ylabel("R$")
    ax.spines[["top", "right"]].set_visible(False)
    # Escreve o valor em cima de cada barra (ex: "12,3 mil")
    ax.bar_label(barras, labels=[f"{v / 1000:,.1f} mil".replace(".", ",")
                                 for v in por_mes["Faturamento"]], fontsize=9)
    fig.tight_layout()
    fig.savefig(arquivo_png, dpi=120)
    plt.close(fig)


# ---------- 4. Salvar o Excel formatado ----------

def salvar_excel(resumos: dict, arquivo_png: Path, arquivo_xlsx: Path) -> None:
    """Escreve cada resumo numa aba e deixa tudo bonito."""
    with pd.ExcelWriter(arquivo_xlsx, engine="openpyxl") as writer:
        for aba, tabela in resumos.items():
            tabela.to_excel(writer, sheet_name=aba, index=False)

    # Agora reabrimos o arquivo para aplicar a formatacao
    planilha = load_workbook(arquivo_xlsx)
    for aba in planilha.worksheets:
        # Cabecalho azul com letra branca
        for celula in aba[1]:
            celula.font = Font(bold=True, color="FFFFFF")
            celula.fill = PatternFill("solid", fgColor=COR_CABECALHO)
            celula.alignment = Alignment(horizontal="center")
        # Coluna "Faturamento" (ou valores em dinheiro) no formato R$
        for coluna in aba.iter_cols(min_row=1):
            if coluna[0].value == "Faturamento":
                for celula in coluna[1:]:
                    celula.number_format = FORMATO_REAIS
        # Largura das colunas de acordo com o texto mais longo
        for coluna in aba.columns:
            maior = max(len(str(c.value)) for c in coluna if c.value is not None)
            aba.column_dimensions[coluna[0].column_letter].width = maior + 4

    # Na aba Resumo, valores em dinheiro (linhas 2 e 4) em R$
    resumo = planilha["Resumo"]
    for linha in (2, 4):
        resumo.cell(row=linha, column=2).number_format = FORMATO_REAIS
    # Grafico colocado na aba Resumo, ao lado dos indicadores
    resumo.add_image(Image(arquivo_png), "D2")

    planilha.save(arquivo_xlsx)


def main():
    parser = argparse.ArgumentParser(description="Gera relatorio de vendas a partir de uma planilha Excel.")
    parser.add_argument("planilha", type=Path, help="arquivo .xlsx com as vendas")
    parser.add_argument("--saida", type=Path, default=Path("saida"), help="pasta do resultado (padrao: saida)")
    args = parser.parse_args()

    args.saida.mkdir(exist_ok=True)
    arquivo_png = args.saida / "grafico_mensal.png"
    arquivo_xlsx = args.saida / "relatorio_vendas.xlsx"

    vendas = ler_vendas(args.planilha)
    resumos = calcular_resumos(vendas)
    gerar_grafico(resumos["Por mes"], arquivo_png)
    salvar_excel(resumos, arquivo_png, arquivo_xlsx)

    total = resumos["Resumo"].iloc[0]["Valor"]
    # Formato brasileiro: 1.234,56 (troca ponto e virgula do formato americano)
    total_br = f"{total:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    print(f"Pronto! {len(vendas)} vendas lidas, faturamento total R$ {total_br}")
    print(f"Relatorio: {arquivo_xlsx}")
    print(f"Grafico:   {arquivo_png}")


if __name__ == "__main__":
    main()
